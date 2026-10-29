"""Write and reload actual messy CSV/XLSX/SQLite businesses."""
from __future__ import annotations
import csv
import io
import re
import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

TABLES = ('products', 'customers', 'orders', 'order_lines', 'purchases', 'debts')
TABLE_NAMES = {
    'products': ('san_pham', 'Products', 'HangHoa'),
    'customers': ('khach_hang', 'Customers', 'KhachHang'),
    'orders': ('don_hang', 'Orders', 'DonHang'),
    'order_lines': ('chi_tiet_don', 'OrderLines', 'ChiTietDon'),
    'purchases': ('phieu_nhap', 'Purchases', 'NhapKho'),
    'debts': ('cong_no', 'Debts', 'CongNo'),
}
COLUMNS = {
    'ma_sp': ('ma_sp', 'SKU', 'MaSP'), 'ten_sp': ('ten_sp', 'ProductName', 'TenHang'),
    'danh_muc': ('danh_muc', 'Category', 'NhomHang'), 'don_vi': ('don_vi', 'Unit', 'DVT'),
    'gia_ban': ('gia_ban', 'Price', 'GiaBan'), 'gia_von': ('gia_von', 'Cost', 'GiaVon'),
    'ma_kh': ('ma_kh', 'CustomerID', 'MaKH'), 'ten_kh': ('ten_kh', 'CustomerName', 'TenKH'),
    'sdt': ('sdt', 'Phone', 'SDT'), 'khu_vuc': ('khu_vuc', 'Region', 'KhuVuc'),
    'ma_don': ('ma_don', 'OrderID', 'MaDon'), 'ngay': ('ngay', 'OrderDate', 'NgayBan'),
    'trang_thai': ('trang_thai', 'Status', 'TrangThai'),
    'voucher_don': ('voucher_don', 'Voucher', 'VoucherDon'),
    'phi_ship': ('phi_ship', 'Shipping', 'PhiShip'), 'vat': ('vat', 'VAT', 'ThueVAT'),
    'dong_id': ('dong_id', 'LineID', 'MaDong'), 'so_luong': ('so_luong', 'Qty', 'SL'),
    'don_gia': ('don_gia', 'UnitPrice', 'DonGia'), 'giam_gia_dong': ('giam_gia_dong', 'LineDiscount', 'GiamGia'),
    'tra_hang': ('tra_hang', 'Returned', 'TraHang'),
    'ten_hien_thi': ('ten_hien_thi', 'DisplayName', 'TenHienThi'),
    'lo_id': ('lo_id', 'LotID', 'MaLo'), 'ngay_nhap': ('ngay_nhap', 'PurchaseDate', 'NgayNhap'),
    'gia_nhap': ('gia_nhap', 'PurchaseCost', 'GiaNhap'), 'han_dung': ('han_dung', 'ExpiryDate', 'HanDung'),
    'so_tien_no': ('so_tien_no', 'AmountDue', 'SoTienNo'), 'ngay_no': ('ngay_no', 'DebtDate', 'NgayNo'),
}


def layout(business_id: str) -> int:
    idx = int(business_id.split('_')[-1]) - 1
    return (idx + (1 if business_id.startswith('dev') else 0)) % 3


def table_name(canonical: str, style: int) -> str:
    return TABLE_NAMES[canonical][style]


def column_name(canonical: str, style: int) -> str:
    return COLUMNS[canonical][style]


def formatted_tables(tables: dict, style: int) -> dict:
    return {
        table_name(key, style): df.rename(columns={c: column_name(c, style) for c in df.columns})
        for key, df in tables.items()
    }


def stable_xlsx(path: Path) -> None:
    """Normalize ZIP member timestamps, removing openpyxl's wall-clock metadata."""
    raw = path.read_bytes()
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw), 'r') as src, zipfile.ZipFile(result, 'w') as dst:
        for item in src.infolist():
            fixed = zipfile.ZipInfo(item.filename, date_time=(2000, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o600 << 16
            fixed.create_system = 3
            payload = src.read(item.filename)
            if item.filename == 'docProps/core.xml':
                payload = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',
                                 rb'\g<1>2000-01-01T00:00:00Z\g<2>', payload)
            dst.writestr(fixed, payload, compress_type=zipfile.ZIP_DEFLATED, compresslevel=6)
    path.write_bytes(result.getvalue())


def write_business(directory: Path, business_id: str, fmt: str, tables: dict, label: str, quirks: list[str]) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    style = layout(business_id)
    named = formatted_tables(tables, style)
    if fmt == 'csv':
        for name, df in named.items():
            df.to_csv(directory / f'{name}.csv', index=False, encoding='utf-8', lineterminator='\n', na_rep='')
    elif fmt == 'xlsx':
        book = Workbook()
        book.remove(book.active)
        for name, df in named.items():
            ws = book.create_sheet(name)
            offset = 3 if name == table_name('orders', style) else 1
            if offset > 1:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(6, len(df.columns)))
                ws.cell(1, 1, f'BÁO CÁO GIAO DỊCH — {business_id}')
                ws.cell(1, 1).font = Font(bold=True, size=13, color='234B67')
            for j, value in enumerate(df.columns, 1):
                cell = ws.cell(offset, j, value)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='276C80')
            for row in df.itertuples(index=False, name=None):
                ws.append(list(row))
            # append() begins at next free row; when header row == 3, this is row 4.
            ws.freeze_panes = f'A{offset+1}'
            for column in ws.columns:
                c = column[offset-1]
                ws.column_dimensions[c.column_letter].width = min(29, max(13, len(str(c.value)) + 3))
        book.properties.created = datetime(2000, 1, 1, 0, 0, 0)
        book.properties.modified = datetime(2000, 1, 1, 0, 0, 0)
        path = directory / 'so_lieu.xlsx'
        book.save(path)
        stable_xlsx(path)
    elif fmt == 'sqlite':
        path = directory / 'kinh_doanh.sqlite'
        with sqlite3.connect(path) as con:
            for name, df in named.items():
                df.to_sql(name, con, index=False, if_exists='replace', chunksize=500, method='multi')
            con.execute('VACUUM')
    else:
        raise ValueError(fmt)
    names = ', '.join(f'`{table_name(x, style)}`' for x in tables)
    cols = ', '.join(f'`{column_name(x, style)}` = `{x}`' for x in ('ma_don', 'ngay', 'ma_sp', 'so_luong', 'don_gia'))
    (directory / 'README.md').write_text(
        f'# {label} — {business_id}\n\n'
        f'Dữ liệu giả lập, không có cá nhân thật. Định dạng: **{fmt.upper()}**. '
        f'Bảng: {names}.\n\n'
        f'Các trường giao dịch thường gặp: {cols}. '
        'Bảng sản phẩm, khách hàng, hóa đơn, chi tiết, phiếu nhập và công nợ liên kết qua mã.\n\n'
        f'**Đặc điểm:** {"; ".join(quirks)}.\n\n'
        'Cần đọc `../../rulebook.md` trước khi tính chỉ tiêu. Dòng trùng `dong_id` là lỗi xuất dữ liệu.\n', encoding='utf-8')


def read_business(directory: Path) -> dict[str, pd.DataFrame]:
    style = layout(directory.name)
    excel = directory / 'so_lieu.xlsx'
    db = directory / 'kinh_doanh.sqlite'
    result = {}
    available_sheets = set(load_workbook(excel, read_only=True).sheetnames) if excel.exists() else set()
    if db.exists():
        with sqlite3.connect(db) as con:
            available_sql = {x[0] for x in con.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
    else:
        available_sql = set()
    for canonical in TABLES:
        name = table_name(canonical, style)
        if excel.exists() and name in available_sheets:
            header_row = 2 if canonical == 'orders' else 0
            df = pd.read_excel(excel, sheet_name=name, header=header_row, dtype=str, keep_default_na=False, engine='openpyxl')
        elif db.exists() and name in available_sql:
            with sqlite3.connect(db) as con:
                df = pd.read_sql_query(f'SELECT * FROM "{name}"', con)
        elif (directory / f'{name}.csv').exists():
            df = pd.read_csv(directory / f'{name}.csv', dtype=str, keep_default_na=False, encoding='utf-8')
        else:
            continue
        inverse = {column_name(x, style): x for x in COLUMNS}
        df = df.rename(columns=inverse)
        result[canonical] = df
    return result
